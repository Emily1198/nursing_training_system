import sqlite3
import pandas as pd
from flask import Flask, render_template_string, request, redirect, url_for, send_file, flash, session
from functools import wraps
import io
from datetime import datetime

import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

app = Flask(__name__)
app.secret_key = 'nursing_secret_key_secure_2026'
DB_NAME = 'nursing_training.db'

# 初始化資料庫 (自動相容舊版結構)
def init_db():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    
    c.execute('''CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    username TEXT UNIQUE NOT NULL,
                    password TEXT NOT NULL
                )''')
    c.execute("INSERT OR IGNORE INTO users (username, password) VALUES ('admin', 'admin123')")
    
    c.execute('''CREATE TABLE IF NOT EXISTS employees (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    emp_no TEXT UNIQUE,
                    name TEXT NOT NULL,
                    position TEXT NOT NULL,
                    status TEXT DEFAULT '在職'
                )''')
    
    c.execute('''CREATE TABLE IF NOT EXISTS courses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    course_date TEXT,
                    title TEXT,
                    category TEXT,
                    hours REAL,
                    source TEXT,
                    method TEXT,
                    has_nursing_points INTEGER DEFAULT 0,
                    has_ltc_points INTEGER DEFAULT 0
                )''')
                
    # 自動補欄位機制 (若舊資料庫缺欄位時使用)
    c.execute("PRAGMA table_info(courses)")
    columns = [col[1] for col in c.fetchall()]
    if 'has_nursing_points' not in columns:
        c.execute("ALTER TABLE courses ADD COLUMN has_nursing_points INTEGER DEFAULT 0")
    if 'has_ltc_points' not in columns:
        c.execute("ALTER TABLE courses ADD COLUMN has_ltc_points INTEGER DEFAULT 0")

    c.execute('''CREATE TABLE IF NOT EXISTS training_records (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    course_id INTEGER,
                    employee_id INTEGER,
                    completion_date TEXT,
                    UNIQUE(course_id, employee_id)
                )''')

    c.execute('''CREATE TABLE IF NOT EXISTS dropdown_options (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    category_type TEXT NOT NULL,
                    option_name TEXT NOT NULL,
                    UNIQUE(category_type, option_name)
                )''')
    
    default_options = [
        ('position', '院長/護理師'), ('position', '主任/護理師'), ('position', '護理師'),
        ('position', '照服員'), ('position', '社工師'), ('position', '營養師'), ('position', '其他'),
        ('course_category', '傳染病與群聚事件'), ('course_category', '感染管制'),
        ('course_category', '消防安全'), ('course_category', '緊急應變'), ('course_category', '性別平等'), ('course_category', '其他'),
        ('source', '長照平台'), ('source', '內訓'), ('source', '外訓'),
        ('method', '線上'), ('method', '實體')
    ]
    c.executemany("INSERT OR IGNORE INTO dropdown_options (category_type, option_name) VALUES (?, ?)", default_options)

    conn.commit()
    conn.close()

init_db()

def login_required(f):
    @wraps(f)
    def decorated_function(*args, **kwargs):
        if 'logged_in' not in session:
            flash("🔒 請先登入系統！")
            return redirect(url_for('login'))
        return f(*args, **kwargs)
    return decorated_function

LOGIN_TEMPLATE = '''
<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <title>系統登入 - 護理之家教育訓練管理系統</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light d-flex align-items-center justify-content-center" style="height: 100vh;">
<div class="card shadow-lg" style="width: 380px;">
    <div class="card-header bg-dark text-white text-center py-3">
        <h4 class="mb-0">🏥 管理員登入</h4>
    </div>
    <div class="card-body p-4">
        {% with messages = get_flashed_messages() %}
          {% if messages %}
            {% for message in messages %}
              <div class="alert alert-danger py-2 small">{{ message|safe }}</div>
            {% endfor %}
          {% endif %}
        {% endwith %}
        <form action="/login" method="post">
            <div class="mb-3">
                <label class="form-label fw-bold">帳號</label>
                <input type="text" name="username" class="form-control" placeholder="預設: admin" required>
            </div>
            <div class="mb-3">
                <label class="form-label fw-bold">密碼</label>
                <input type="password" name="password" class="form-control" placeholder="預設: admin123" required>
            </div>
            <button type="submit" class="btn btn-primary w-100 fw-bold py-2">登入系統</button>
        </form>
    </div>
</div>
</body>
</html>
'''

HTML_TEMPLATE = '''
<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <title>護理之家教育訓練統計管理系統</title>
    <link href="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/css/bootstrap.min.css" rel="stylesheet">
</head>
<body class="bg-light pb-5">
<nav class="navbar navbar-expand-lg navbar-dark bg-dark mb-4 shadow">
    <div class="container-fluid">
        <span class="navbar-brand mb-0 h1">🏥 護理之家教育訓練統計管理系統</span>
        <div class="d-flex align-items-center">
            <span class="text-light me-3 small">👤 目前使用者：<strong>{{ session['username'] }}</strong></span>
            <button class="btn btn-sm btn-outline-info me-2" data-bs-toggle="modal" data-bs-target="#optionsModal">⚙️ 下拉選單管理</button>
            <button class="btn btn-sm btn-outline-warning me-2" data-bs-toggle="modal" data-bs-target="#accountModal">🔒 帳號/密碼設定</button>
            <a href="/logout" class="btn btn-sm btn-outline-light">🚪 登出</a>
        </div>
    </div>
</nav>

<div class="container">
    {% with messages = get_flashed_messages() %}
      {% if messages %}
        {% for message in messages %}
          <div class="alert alert-warning alert-dismissible fade show shadow-sm" role="alert">
            {{ message|safe }}
            <button type="button" class="btn-close" data-bs-dismiss="alert"></button>
          </div>
        {% endfor %}
      {% endif %}
    {% endwith %}

    <div class="row">
        <!-- 1. 新增員工 -->
        <div class="col-md-5 mb-4">
            <div class="card shadow-sm h-100">
                <div class="card-header bg-primary text-white font-weight-bold">👤 新增員工</div>
                <div class="card-body">
                    <form action="/add_employee" method="post">
                        <div class="mb-2"><input type="text" name="emp_no" class="form-control" placeholder="工號 (如: EMP001)" required></div>
                        <div class="mb-2"><input type="text" name="name" class="form-control" placeholder="姓名" required></div>
                        <div class="mb-2">
                            <select name="position" class="form-select" required>
                                <option value="">選擇職位...</option>
                                {% for opt in opts.position %}
                                <option value="{{ opt }}">{{ opt }}</option>
                                {% endfor %}
                            </select>
                        </div>
                        <button type="submit" class="btn btn-primary w-100 mt-2">儲存員工資料</button>
                    </form>
                </div>
            </div>
        </div>

        <!-- 2. 建立課程 (新增積分選單) -->
        <div class="col-md-7 mb-4">
            <div class="card shadow-sm h-100">
                <div class="card-header bg-success text-white font-weight-bold">📚 建立新課程</div>
                <div class="card-body">
                    <form action="/add_course" method="post">
                        <div class="row g-2 mb-2">
                            <div class="col-md-4"><label class="form-label mb-0 small fw-bold">課程日期</label><input type="date" name="course_date" class="form-control" required></div>
                            <div class="col-md-5"><label class="form-label mb-0 small fw-bold">課程名稱</label><input type="text" name="title" class="form-control" placeholder="如: 傳染病防治概論" required></div>
                            <div class="col-md-3">
                                <label class="form-label mb-0 small fw-bold">課程類別</label>
                                <select name="category" class="form-select" required>
                                    {% for opt in opts.course_category %}
                                    <option value="{{ opt }}">{{ opt }}</option>
                                    {% endfor %}
                                </select>
                            </div>
                        </div>
                        <div class="row g-2 mb-2">
                            <div class="col-md-4"><label class="form-label mb-0 small fw-bold">時數</label><input type="number" step="0.5" name="hours" class="form-control" placeholder="如 1.0" required></div>
                            <div class="col-md-4">
                                <label class="form-label mb-0 small fw-bold">訓練來源</label>
                                <select name="source" class="form-select" required>
                                    {% for opt in opts.source %}
                                    <option value="{{ opt }}">{{ opt }}</option>
                                    {% endfor %}
                                </select>
                            </div>
                            <div class="col-md-4">
                                <label class="form-label mb-0 small fw-bold">上課方式</label>
                                <select name="method" class="form-select" required>
                                    {% for opt in opts.method %}
                                    <option value="{{ opt }}">{{ opt }}</option>
                                    {% endfor %}
                                </select>
                            </div>
                        </div>
                        <!-- 新增積分核取欄位 -->
                        <div class="mb-3 bg-light p-2 rounded border">
                            <span class="small fw-bold me-3">🎖️ 積分屬性設定：</span>
                            <div class="form-check form-check-inline">
                                <input class="form-check-input" type="checkbox" name="has_nursing_points" id="chk_nursing" value="1">
                                <label class="form-check-input-label small fw-bold text-primary" for="chk_nursing">含有「護理積分」</label>
                            </div>
                            <div class="form-check form-check-inline">
                                <input class="form-check-input" type="checkbox" name="has_ltc_points" id="chk_ltc" value="1">
                                <label class="form-check-input-label small fw-bold text-success" for="chk_ltc">含有「長照積分」</label>
                            </div>
                        </div>
                        <button type="submit" class="btn btn-success w-100 fw-bold">建立課程</button>
                    </form>
                </div>
            </div>
        </div>
    </div>

    <!-- 3. 登錄完訓紀錄 -->
    <div class="card shadow-sm mb-4 border-info">
        <div class="card-header bg-info text-dark fw-bold">✍️ 登錄完訓紀錄</div>
        <div class="card-body bg-white">
            <ul class="nav nav-tabs" id="recordTab" role="tablist">
                <li class="nav-item">
                    <button class="nav-link active fw-bold" id="single-tab" data-bs-toggle="tab" data-bs-target="#single-panel" type="button">方式 A：單人手動 Key 入</button>
                </li>
                <li class="nav-item">
                    <button class="nav-link fw-bold" id="batch-tab" data-bs-toggle="tab" data-bs-target="#batch-panel" type="button">方式 B：整批 Excel 匯入</button>
                </li>
            </ul>
            <div class="tab-content pt-3">
                <div class="tab-pane fade show active" id="single-panel">
                    <form action="/add_single_record" method="post" class="row g-2 align-items-end">
                        <div class="col-md-4">
                            <label class="form-label small fw-bold">選擇課程</label>
                            <select name="course_id" class="form-select" required>
                                <option value="">請選擇課程...</option>
                                {% for c in courses %}
                                <option value="{{ c[0] }}">{{ c[1] }} - {{ c[2] }} ({{ c[3] }})</option>
                                {% endfor %}
                            </select>
                        </div>
                        <div class="col-md-3">
                            <label class="form-label small fw-bold">選擇完訓員工</label>
                            <select name="employee_id" class="form-select" required>
                                <option value="">請選擇員工...</option>
                                {% for e in employees %}
                                <option value="{{ e[0] }}">{{ e[1] }} - {{ e[2] }} ({{ e[3] }})</option>
                                {% endfor %}
                            </select>
                        </div>
                        <div class="col-md-3">
                            <label class="form-label small fw-bold">實際完訓日期</label>
                            <input type="date" name="completion_date" class="form-control" required>
                        </div>
                        <div class="col-md-2">
                            <button type="submit" class="btn btn-info w-100 text-dark fw-bold">新增完訓</button>
                        </div>
                    </form>
                </div>
                <div class="tab-pane fade" id="batch-panel">
                    <form action="/upload_batch_records" method="post" enctype="multipart/form-data" class="row g-2 align-items-end">
                        <div class="col-md-5">
                            <label class="form-label small fw-bold">選擇課程</label>
                            <select name="course_id" class="form-select" required>
                                <option value="">請選擇課程...</option>
                                {% for c in courses %}
                                <option value="{{ c[0] }}">{{ c[1] }} - {{ c[2] }} ({{ c[3] }})</option>
                                {% endfor %}
                            </select>
                        </div>
                        <div class="col-md-5">
                            <label class="form-label small fw-bold">上傳完訓名單 Excel (含「姓名」與可選「完訓日期」)</label>
                            <input type="file" name="excel_file" class="form-control" accept=".xlsx, .xls" required>
                        </div>
                        <div class="col-md-2">
                            <button type="submit" class="btn btn-info w-100 text-dark fw-bold">整批匯入</button>
                        </div>
                    </form>
                </div>
            </div>
        </div>
    </div>

    <!-- 4. 表格維護區 -->
    <div class="row">
        <div class="col-md-5 mb-4">
            <div class="card shadow-sm h-100">
                <div class="card-header bg-secondary text-white fw-bold">👥 目前員工名冊 (共 {{ employees|length }} 人)</div>
                <div class="card-body p-0" style="max-height: 300px; overflow-y: auto;">
                    <table class="table table-striped table-hover mb-0 small align-middle">
                        <thead class="table-light sticky-top">
                            <tr><th>工號</th><th>姓名</th><th>職位</th><th>狀態</th><th>操作</th></tr>
                        </thead>
                        <tbody>
                            {% for e in employees %}
                            <tr>
                                <form action="/update_employee/{{ e[0] }}" method="post">
                                    <td>{{ e[1] }}</td>
                                    <td class="fw-bold">{{ e[2] }}</td>
                                    <td>{{ e[3] }}</td>
                                    <td>
                                        <select name="status" class="form-select form-select-sm p-1">
                                            <option value="在職" {% if e[4] == '在職' %}selected{% endif %}>在職</option>
                                            <option value="離職" {% if e[4] == '離職' %}selected{% endif %}>離職</option>
                                        </select>
                                    </td>
                                    <td><button type="submit" class="btn btn-outline-primary btn-sm py-0 px-1">更新</button></td>
                                </form>
                            </tr>
                            {% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
        </div>

        <div class="col-md-7 mb-4">
            <div class="card shadow-sm h-100">
                <div class="card-header bg-secondary text-white fw-bold">📖 目前課程清單 (共 {{ courses|length }} 門)</div>
                <div class="card-body p-0" style="max-height: 300px; overflow-y: auto;">
                    <table class="table table-striped table-hover mb-0 small align-middle">
                        <thead class="table-light sticky-top">
                            <tr><th>日期</th><th>課程名稱</th><th>類別</th><th>時數</th><th>積分屬性</th><th>完訓</th></tr>
                        </thead>
                        <tbody>
                            {% for c in courses %}
                            <tr>
                                <td>{{ c[1] }}</td>
                                <td class="fw-bold">{{ c[2] }}</td>
                                <td>{{ c[3] }}</td>
                                <td>{{ c[4] }}</td>
                                <td>
                                    {% if c[7] %}<span class="badge bg-primary">護理</span>{% endif %}
                                    {% if c[8] %}<span class="badge bg-success">長照</span>{% endif %}
                                    {% if not c[7] and not c[8] %}<span class="text-muted small">無</span>{% endif %}
                                </td>
                                <td><span class="badge bg-info text-dark">{{ c[9] }} 人</span></td>
                            </tr>
                            {% endfor %}
                        </tbody>
                    </table>
                </div>
            </div>
        </div>
    </div>

    <!-- 5. 🌟 查核統計與一鍵產生 Excel 總報表 (含積分篩選) 🌟 -->
    <div class="card shadow-sm mb-4 border-dark">
        <div class="card-header bg-dark text-white font-weight-bold">📊 查核統計與一鍵產生 Excel 總報表</div>
        <div class="card-body bg-white">
            <form action="/export_excel_router" method="get" class="row g-3">
                <div class="col-md-6">
                    <label class="form-label fw-bold text-primary">📄 請選擇欲產生的報表格式</label>
                    <select name="report_type" class="form-select border-primary fw-bold">
                        <option value="cross_tab">📋 方案 A：評鑑專用「交叉查核矩陣總表」 (日期/X比對表)</option>
                        <option value="multi_sheet">📊 方案 B：簡易統計報表 (含多個統計分頁 Sheets)</option>
                    </select>
                </div>
                <div class="col-md-6">
                    <label class="form-label fw-bold">課程類別主題 (篩選)</label>
                    <select name="category" class="form-select">
                        <option value="">全部類別 (預設)</option>
                        {% for opt in opts.course_category %}
                        <option value="{{ opt }}">{{ opt }}</option>
                        {% endfor %}
                    </select>
                </div>
                <!-- 新增積分條件篩選器 -->
                <div class="col-md-4">
                    <label class="form-label fw-bold">🎖️ 積分屬性 (篩選)</label>
                    <select name="points_filter" class="form-select">
                        <option value="">全部課程 (不限積分)</option>
                        <option value="nursing">僅包含「護理積分」課程</option>
                        <option value="ltc">僅包含「長照積分」課程</option>
                        <option value="both">必須同時含「護理與長照積分」</option>
                    </select>
                </div>
                <div class="col-md-4">
                    <label class="form-label fw-bold">開始日期 (完訓/開課)</label>
                    <input type="date" name="start_date" class="form-control">
                </div>
                <div class="col-md-4">
                    <label class="form-label fw-bold">結束日期 (完訓/開課)</label>
                    <input type="date" name="end_date" class="form-control">
                </div>
                <div class="col-12 text-end mt-3">
                    <button type="submit" class="btn btn-danger btn-lg fw-bold w-100">
                        🚀 一鍵篩選並匯出 Excel 報表
                    </button>
                </div>
            </form>
        </div>
    </div>
</div>

<!-- Modal 1: 下拉選單管理 -->
<div class="modal fade" id="optionsModal" tabindex="-1">
    <div class="modal-dialog modal-lg">
        <div class="modal-content">
            <div class="modal-header bg-info text-dark">
                <h5 class="modal-title fw-bold">⚙️ 管理下拉選單選項</h5>
                <button type="button" class="btn-close" data-bs-dismiss="modal"></button>
            </div>
            <div class="modal-body">
                <form action="/add_dropdown_option" method="post" class="row g-2 mb-4 bg-light p-3 rounded border">
                    <div class="col-md-4">
                        <label class="form-label fw-bold small">選擇類別</label>
                        <select name="category_type" class="form-select" required>
                            <option value="position">員工職位</option>
                            <option value="course_category">課程類別</option>
                            <option value="source">訓練來源</option>
                            <option value="method">上課方式</option>
                        </select>
                    </div>
                    <div class="col-md-5">
                        <label class="form-label fw-bold small">新增選項名稱</label>
                        <input type="text" name="option_name" class="form-control" placeholder="例如：物理治療師、線上直播" required>
                    </div>
                    <div class="col-md-3 d-flex align-items-end">
                        <button type="submit" class="btn btn-success w-100 fw-bold">➕ 新增選項</button>
                    </div>
                </form>

                <h6 class="fw-bold mb-2">現有下拉選項一覽：</h6>
                <div class="row">
                    {% for cat_type, cat_title in [('position','員工職位'), ('course_category','課程類別'), ('source','訓練來源'), ('method','上課方式')] %}
                    <div class="col-md-6 mb-3">
                        <div class="card">
                            <div class="card-header py-1 fw-bold bg-secondary text-white small">{{ cat_title }}</div>
                            <ul class="list-group list-group-flush small" style="max-height: 150px; overflow-y: auto;">
                                {% for opt in opts[cat_type] %}
                                <li class="list-group-item d-flex justify-content-between align-items-center py-1">
                                    {{ opt }}
                                    <form action="/delete_dropdown_option" method="post" style="display:inline;">
                                        <input type="hidden" name="category_type" value="{{ cat_type }}">
                                        <input type="hidden" name="option_name" value="{{ opt }}">
                                        <button type="submit" class="btn btn-outline-danger btn-sm py-0 px-1" onclick="return confirm('確定刪除此選項？')">刪除</button>
                                    </form>
                                </li>
                                {% endfor %}
                            </ul>
                        </div>
                    </div>
                    {% endfor %}
                </div>
            </div>
            <div class="modal-footer">
                <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">關閉</button>
            </div>
        </div>
    </div>
</div>

<!-- Modal 2: 帳號密碼修改 -->
<div class="modal fade" id="accountModal" tabindex="-1">
    <div class="modal-dialog">
        <div class="modal-content">
            <div class="modal-header bg-warning">
                <h5 class="modal-title fw-bold">🔒 帳號與密碼設定</h5>
                <button type="button" class="btn-close" data-bs-dismiss="modal"></button>
            </div>
            <form action="/update_account" method="post">
                <div class="modal-body">
                    <div class="mb-3">
                        <label class="form-label fw-bold">帳號名稱</label>
                        <input type="text" name="new_username" class="form-control" value="{{ session['username'] }}" required>
                    </div>
                    <div class="mb-3">
                        <label class="form-label fw-bold">目前舊密碼 <span class="text-danger">*</span></label>
                        <input type="password" name="old_password" class="form-control" placeholder="請輸入舊密碼" required>
                    </div>
                    <hr>
                    <div class="mb-3">
                        <label class="form-label fw-bold">新密碼 <small class="text-muted">(不修改請留空)</small></label>
                        <input type="password" name="new_password" class="form-control" placeholder="留空代表維持原密碼">
                    </div>
                </div>
                <div class="modal-footer">
                    <button type="button" class="btn btn-secondary" data-bs-dismiss="modal">取消</button>
                    <button type="submit" class="btn btn-warning fw-bold">儲存變更</button>
                </div>
            </form>
        </div>
    </div>
</div>

<script src="https://cdn.jsdelivr.net/npm/bootstrap@5.3.0/dist/js/bootstrap.bundle.min.js"></script>
</body>
</html>
'''

def get_dropdown_options():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("SELECT category_type, option_name FROM dropdown_options")
    rows = c.fetchall()
    conn.close()
    
    opts = {'position': [], 'course_category': [], 'source': [], 'method': []}
    for cat, name in rows:
        if cat in opts:
            opts[cat].append(name)
    return opts

@app.route('/login', methods=['GET', 'POST'])
def login():
    if request.method == 'POST':
        username = request.form['username'].strip()
        password = request.form['password'].strip()
        conn = sqlite3.connect(DB_NAME)
        c = conn.cursor()
        c.execute("SELECT id, username FROM users WHERE username = ? AND password = ?", (username, password))
        user = c.fetchone()
        conn.close()
        if user:
            session['logged_in'] = True
            session['user_id'] = user[0]
            session['username'] = user[1]
            return redirect(url_for('index'))
        else:
            flash("❌ 帳號或密碼錯誤！")
            return redirect(url_for('login'))
    return render_template_string(LOGIN_TEMPLATE)

@app.route('/logout')
def logout():
    session.clear()
    flash("👋 已成功登出系統。")
    return redirect(url_for('login'))

@app.route('/')
@login_required
def index():
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("SELECT id, emp_no, name, position, status FROM employees ORDER BY emp_no ASC")
    employees = c.fetchall()
    
    c.execute('''
        SELECT c.id, c.course_date, c.title, c.category, c.hours, c.source, c.method,
               c.has_nursing_points, c.has_ltc_points,
               COUNT(r.id) as attendee_count
        FROM courses c
        LEFT JOIN training_records r ON c.id = r.course_id
        GROUP BY c.id
        ORDER BY c.course_date DESC
    ''')
    courses = c.fetchall()
    conn.close()
    
    opts = get_dropdown_options()
    return render_template_string(HTML_TEMPLATE, employees=employees, courses=courses, opts=opts)

@app.route('/add_dropdown_option', methods=['POST'])
@login_required
def add_dropdown_option():
    cat_type = request.form['category_type']
    opt_name = request.form['option_name'].strip()
    
    if opt_name:
        conn = sqlite3.connect(DB_NAME)
        c = conn.cursor()
        try:
            c.execute("INSERT INTO dropdown_options (category_type, option_name) VALUES (?, ?)", (cat_type, opt_name))
            conn.commit()
            flash(f"✅ 已成功新增選項【{opt_name}】！")
        except sqlite3.IntegrityError:
            flash(f"⚠️ 選項【{opt_name}】已經存在！")
        finally:
            conn.close()
    return redirect(url_for('index'))

@app.route('/delete_dropdown_option', methods=['POST'])
@login_required
def delete_dropdown_option():
    cat_type = request.form['category_type']
    opt_name = request.form['option_name']
    
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("DELETE FROM dropdown_options WHERE category_type = ? AND option_name = ?", (cat_type, opt_name))
    conn.commit()
    conn.close()
    flash(f"🗑️ 已刪除選項【{opt_name}】。")
    return redirect(url_for('index'))

@app.route('/update_account', methods=['POST'])
@login_required
def update_account():
    new_username = request.form['new_username'].strip()
    old_password = request.form['old_password'].strip()
    new_password = request.form['new_password'].strip()
    user_id = session.get('user_id')

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("SELECT password FROM users WHERE id = ?", (user_id,))
    user = c.fetchone()
    
    if not user or user[0] != old_password:
        flash("❌ 舊密碼輸入錯誤！")
    else:
        c.execute("SELECT id FROM users WHERE username = ? AND id != ?", (new_username, user_id))
        if c.fetchone():
            flash(f"❌ 帳號名稱【{new_username}】已被使用！")
        else:
            final_password = new_password if new_password else old_password
            c.execute("UPDATE users SET username = ?, password = ? WHERE id = ?", (new_username, final_password, user_id))
            conn.commit()
            session['username'] = new_username
            flash("✅ 帳號/密碼更新成功！")
            
    conn.close()
    return redirect(url_for('index'))

@app.route('/add_employee', methods=['POST'])
@login_required
def add_employee():
    emp_no = request.form['emp_no'].strip()
    name = request.form['name'].strip()
    position = request.form['position']
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    try:
        c.execute("INSERT INTO employees (emp_no, name, position) VALUES (?, ?, ?)", (emp_no, name, position))
        conn.commit()
        flash(f"✅ 已新增員工：{name} ({position})")
    except sqlite3.IntegrityError:
        flash("❌ 新增失敗：工號已存在！")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/update_employee/<int:emp_id>', methods=['POST'])
@login_required
def update_employee(emp_id):
    status = request.form['status']
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("UPDATE employees SET status = ? WHERE id = ?", (status, emp_id))
    conn.commit()
    conn.close()
    flash("✅ 員工狀態更新成功！")
    return redirect(url_for('index'))

@app.route('/add_course', methods=['POST'])
@login_required
def add_course():
    c_date = request.form['course_date']
    title = request.form['title'].strip()
    category = request.form['category']
    hours = float(request.form['hours'])
    source = request.form['source']
    method = request.form['method']
    has_nursing = 1 if request.form.get('has_nursing_points') else 0
    has_ltc = 1 if request.form.get('has_ltc_points') else 0
    
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("INSERT INTO courses (course_date, title, category, hours, source, method, has_nursing_points, has_ltc_points) VALUES (?,?,?,?,?,?,?,?)",
              (c_date, title, category, hours, source, method, has_nursing, has_ltc))
    conn.commit()
    conn.close()
    flash(f"✅ 課程【{title}】建立成功！")
    return redirect(url_for('index'))

@app.route('/add_single_record', methods=['POST'])
@login_required
def add_single_record():
    course_id = request.form['course_id']
    employee_id = request.form['employee_id']
    completion_date = request.form['completion_date']
    
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    try:
        c.execute("INSERT INTO training_records (course_id, employee_id, completion_date) VALUES (?, ?, ?)",
                  (course_id, employee_id, completion_date))
        conn.commit()
        flash("✅ 完訓紀錄登錄成功！")
    except sqlite3.IntegrityError:
        flash("⚠️ 該員工已有此課程之完訓紀錄。")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/upload_batch_records', methods=['POST'])
@login_required
def upload_batch_records():
    course_id = request.form['course_id']
    file = request.files.get('excel_file')
    
    if not file or file.filename == '':
        flash("❌ 請選擇 Excel 檔案！")
        return redirect(url_for('index'))
        
    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()
    c.execute("SELECT course_date FROM courses WHERE id = ?", (course_id,))
    course_info = c.fetchone()
    default_date = course_info[0] if course_info else ''
    
    try:
        df = pd.read_excel(file)
        if '姓名' not in df.columns:
            flash("❌ Excel 檔案內缺少「姓名」欄位！")
            conn.close()
            return redirect(url_for('index'))
            
        matched = 0
        for idx, row in df.iterrows():
            name = str(row['姓名']).strip()
            c_date = str(row['完訓日期']).split(' ')[0].strip() if '完訓日期' in df.columns and pd.notna(row['完訓日期']) else default_date
            
            c.execute("SELECT id FROM employees WHERE name = ?", (name,))
            emp = c.fetchone()
            if emp:
                try:
                    c.execute("INSERT INTO training_records (course_id, employee_id, completion_date) VALUES (?, ?, ?)",
                              (course_id, emp[0], c_date))
                    matched += 1
                except sqlite3.IntegrityError:
                    pass
        conn.commit()
        flash(f"✅ 批次匯入完成！成功登錄 {matched} 人完訓。")
    except Exception as e:
        flash(f"❌ Excel 處理失敗：{str(e)}")
    finally:
        conn.close()
    return redirect(url_for('index'))

@app.route('/export_excel_router', methods=['GET'])
@login_required
def export_excel_router():
    report_type = request.args.get('report_type', 'cross_tab')
    
    if report_type == 'multi_sheet':
        return export_multi_sheet_excel()
    else:
        return export_cross_tab_excel()

# -------------------------------------------------------------
# 報表 A：評鑑專用交叉查核矩陣總表 (支援積分標註與篩選)
# -------------------------------------------------------------
def export_cross_tab_excel():
    category_filter = request.args.get('category', '').strip()
    points_filter = request.args.get('points_filter', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()

    conn = sqlite3.connect(DB_NAME)
    c = conn.cursor()

    course_query = "SELECT id, course_date, title, hours, source, method, has_nursing_points, has_ltc_points FROM courses WHERE 1=1"
    params = []
    if category_filter:
        course_query += " AND category = ?"
        params.append(category_filter)
    if points_filter == 'nursing':
        course_query += " AND has_nursing_points = 1"
    elif points_filter == 'ltc':
        course_query += " AND has_ltc_points = 1"
    elif points_filter == 'both':
        course_query += " AND has_nursing_points = 1 AND has_ltc_points = 1"

    if start_date:
        course_query += " AND course_date >= ?"
        params.append(start_date)
    if end_date:
        course_query += " AND course_date <= ?"
        params.append(end_date)
    course_query += " ORDER BY course_date DESC"

    c.execute(course_query, params)
    raw_courses = c.fetchall()

    if not raw_courses:
        flash("⚠️ 指定篩選條件下無任何課程資料！")
        conn.close()
        return redirect(url_for('index'))

    courses_by_year = {}
    for c_id, c_date, title, hours, source, method, h_nursing, h_ltc in raw_courses:
        try:
            dt = datetime.strptime(c_date, '%Y-%m-%d')
            roc_year = str(dt.year - 1911)
        except:
            roc_year = "其他"
            
        if roc_year not in courses_by_year:
            courses_by_year[roc_year] = []
        
        # 抬頭加上積分標註
        tag = ""
        if h_nursing and h_ltc: tag = "[護/長] "
        elif h_nursing: tag = "[護] "
        elif h_ltc: tag = "[長] "
        
        courses_by_year[roc_year].append({
            'id': c_id, 'date': c_date, 'title': tag + title, 
            'hours': hours, 'source': source, 'method': method
        })

    c.execute("SELECT id, position, name FROM employees WHERE status = '在職' ORDER BY id ASC")
    employees = c.fetchall()

    c.execute("SELECT course_id, employee_id, completion_date FROM training_records")
    records = {(r[0], r[1]): r[2] for r in c.fetchall()}

    conn.close()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "教育課程清冊總表"

    font_title = Font(name='標楷體', size=16, bold=True)
    font_subtitle = Font(name='標楷體', size=10)
    font_header = Font(name='新細明體', size=10, bold=True)
    font_data = Font(name='新細明體', size=10)
    
    align_center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    align_right = Alignment(horizontal='right', vertical='center')

    fill_pink = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_orange = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
    border_thin = Side(border_style="thin", color="000000")
    box_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    all_dates = [c['date'] for y in courses_by_year for c in courses_by_year[y]]
    min_date_str = min(all_dates) if all_dates else ""
    max_date_str = max(all_dates) if all_dates else ""
    
    title_text = f"{category_filter if category_filter else '教育課程'} 清冊總表"
    if min_date_str and max_date_str:
        d1 = datetime.strptime(min_date_str, '%Y-%m-%d')
        d2 = datetime.strptime(max_date_str, '%Y-%m-%d')
        title_text = f"{d1.year-1911}/{d1.month}/{d1.day} - {d2.year-1911}年-{title_text}"

    ws.row_dimensions[1].height = 30
    ws.cell(row=1, column=1, value=title_text).font = font_title
    ws.cell(row=1, column=1).alignment = align_center

    ws.row_dimensions[2].height = 18
    ws.cell(row=2, column=1, value="每年安排相關教育訓練(>4Hr/人)，此次檢附教育訓練之課程佐證(簽到、照片、講義)").font = font_subtitle
    
    today_roc = f"{datetime.now().year - 1911}/{datetime.now().month}/{datetime.now().day}"
    
    col_idx = 4
    year_col_spans = {}
    
    for year, c_list in courses_by_year.items():
        start_col = col_idx
        col_idx += len(c_list)
        end_col = col_idx
        year_col_spans[year] = (start_col, end_col, c_list)
        col_idx += 1

    total_max_cols = col_idx - 1
    
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_max_cols)
    ws.cell(row=2, column=total_max_cols, value=f"整理日期：{today_roc}").font = font_subtitle
    ws.cell(row=2, column=total_max_cols).alignment = align_right

    ws.merge_cells('A3:A5'); ws.cell(row=3, column=1, value="序號").alignment = align_center
    ws.merge_cells('B3:B5'); ws.cell(row=3, column=2, value="職稱").alignment = align_center
    ws.merge_cells('C3:C5'); ws.cell(row=3, column=3, value="姓名").alignment = align_center

    for year, (s_col, e_col, c_list) in year_col_spans.items():
        ws.merge_cells(start_row=3, start_column=s_col, end_row=3, end_column=e_col)
        cell_y = ws.cell(row=3, column=s_col, value=f"{year} 年度" if year != "其他" else "其他年度")
        cell_y.alignment = align_center
        cell_y.font = font_header

        for idx, course in enumerate(c_list):
            curr_col = s_col + idx
            c_name_cell = ws.cell(row=4, column=curr_col, value=course['title'])
            c_name_cell.alignment = align_center
            c_name_cell.font = font_data
            
            method_str = f"{course['source']}/{int(course['hours'])}Hr" if course['hours'].is_integer() else f"{course['source']}/{course['hours']}Hr"
            ws.cell(row=5, column=curr_col, value=method_str).alignment = align_center
            
        ws.merge_cells(start_row=4, start_column=e_col, end_row=5, end_column=e_col)
        ws.cell(row=4, column=e_col, value="總時數").alignment = align_center

    current_row = 6
    for emp_idx, (emp_id, pos, name) in enumerate(employees, start=1):
        ws.row_dimensions[current_row].height = 20
        ws.cell(row=current_row, column=1, value=emp_idx).alignment = align_center
        ws.cell(row=current_row, column=2, value=pos).alignment = align_center
        ws.cell(row=current_row, column=3, value=name).alignment = align_center

        for year, (s_col, e_col, c_list) in year_col_spans.items():
            year_hours_sum = 0
            for idx, course in enumerate(c_list):
                curr_col = s_col + idx
                key = (course['id'], emp_id)
                if key in records:
                    comp_date = records[key]
                    try:
                        dt = datetime.strptime(comp_date, '%Y-%m-%d')
                        disp_date = f"{dt.year-1911}.{dt.month:02d}.{dt.day:02d}"
                    except:
                        disp_date = comp_date
                    ws.cell(row=current_row, column=curr_col, value=disp_date).alignment = align_center
                    year_hours_sum += course['hours']
                else:
                    ws.cell(row=current_row, column=curr_col, value="X").alignment = align_center

            ws.cell(row=current_row, column=e_col, value=year_hours_sum if year_hours_sum > 0 else 0).alignment = align_center

        current_row += 1

    for r in range(3, current_row):
        for c_i in range(1, total_max_cols + 1):
            cell = ws.cell(row=r, column=c_i)
            cell.border = box_border
            if r in [3, 4, 5]:
                cell.font = font_header
                cell.fill = fill_pink if c_i <= (total_max_cols // 2 + 1) else fill_orange

    ws.column_dimensions['A'].width = 6
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 10
    for c_i in range(4, total_max_cols + 1):
        ws.column_dimensions[get_column_letter(c_i)].width = 16

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    
    filename = f"護理之家_評鑑專用交叉查核總表_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True)

# -------------------------------------------------------------
# 報表 B：簡易統計報表 (多分頁 Excel，含積分專屬欄位)
# -------------------------------------------------------------
def export_multi_sheet_excel():
    category_filter = request.args.get('category', '').strip()
    points_filter = request.args.get('points_filter', '').strip()
    start_date = request.args.get('start_date', '').strip()
    end_date = request.args.get('end_date', '').strip()

    conn = sqlite3.connect(DB_NAME)
    
    # 1. Sheet1: 員工完訓時數統計總表
    emp_query = '''
        SELECT e.emp_no AS 工號, e.name AS 姓名, e.position AS 職稱, e.status AS 狀態,
               COALESCE(SUM(c.hours), 0) AS 完訓總時數,
               COUNT(r.id) AS 完訓總門數
        FROM employees e
        LEFT JOIN training_records r ON e.id = r.employee_id
        LEFT JOIN courses c ON r.course_id = c.id
        WHERE 1=1
    '''
    params = []
    if category_filter:
        emp_query += " AND (c.category = ? OR c.category IS NULL)"
        params.append(category_filter)
    if points_filter == 'nursing':
        emp_query += " AND (c.has_nursing_points = 1 OR c.id IS NULL)"
    elif points_filter == 'ltc':
        emp_query += " AND (c.has_ltc_points = 1 OR c.id IS NULL)"
    elif points_filter == 'both':
        emp_query += " AND (c.has_nursing_points = 1 AND c.has_ltc_points = 1 OR c.id IS NULL)"

    if start_date:
        emp_query += " AND (c.course_date >= ? OR c.course_date IS NULL)"
        params.append(start_date)
    if end_date:
        emp_query += " AND (c.course_date <= ? OR c.course_date IS NULL)"
        params.append(end_date)
    emp_query += " GROUP BY e.id ORDER BY e.emp_no ASC"
    
    df_emp_summary = pd.read_sql_query(emp_query, conn, params=params)

    # 2. Sheet2: 課程開課統計表 (含積分)
    course_query = '''
        SELECT c.course_date AS 開課日期, c.title AS 課程名稱, c.category AS 課程類別,
               c.hours AS 時數, c.source AS 訓練來源, c.method AS 上課方式,
               CASE WHEN c.has_nursing_points = 1 THEN '是' ELSE '否' END AS 護理積分,
               CASE WHEN c.has_ltc_points = 1 THEN '是' ELSE '否' END AS 長照積分,
               COUNT(r.id) AS 完訓人數
        FROM courses c
        LEFT JOIN training_records r ON c.id = r.course_id
        WHERE 1=1
    '''
    c_params = []
    if category_filter:
        course_query += " AND c.category = ?"
        c_params.append(category_filter)
    if points_filter == 'nursing':
        course_query += " AND c.has_nursing_points = 1"
    elif points_filter == 'ltc':
        course_query += " AND c.has_ltc_points = 1"
    elif points_filter == 'both':
        course_query += " AND c.has_nursing_points = 1 AND c.has_ltc_points = 1"

    if start_date:
        course_query += " AND c.course_date >= ?"
        c_params.append(start_date)
    if end_date:
        course_query += " AND c.course_date <= ?"
        c_params.append(end_date)
    course_query += " GROUP BY c.id ORDER BY c.course_date DESC"

    df_course_summary = pd.read_sql_query(course_query, conn, params=c_params)

    # 3. Sheet3: 完整完訓明細流水帳 (含積分)
    detail_query = '''
        SELECT e.emp_no AS 員工工號, e.name AS 員工姓名, e.position AS 職稱,
               c.course_date AS 課程日期, r.completion_date AS 實際完訓日期,
               c.title AS 課程名稱, c.category AS 課程類別, c.hours AS 時數,
               c.source AS 訓練來源, c.method AS 上課方式,
               CASE WHEN c.has_nursing_points = 1 THEN '是' ELSE '否' END AS 護理積分,
               CASE WHEN c.has_ltc_points = 1 THEN '是' ELSE '否' END AS 長照積分
        FROM training_records r
        JOIN employees e ON r.employee_id = e.id
        JOIN courses c ON r.course_id = c.id
        WHERE 1=1
    '''
    d_params = []
    if category_filter:
        detail_query += " AND c.category = ?"
        d_params.append(category_filter)
    if points_filter == 'nursing':
        detail_query += " AND c.has_nursing_points = 1"
    elif points_filter == 'ltc':
        detail_query += " AND c.has_ltc_points = 1"
    elif points_filter == 'both':
        detail_query += " AND c.has_nursing_points = 1 AND c.has_ltc_points = 1"

    if start_date:
        detail_query += " AND c.course_date >= ?"
        d_params.append(start_date)
    if end_date:
        detail_query += " AND c.course_date <= ?"
        d_params.append(end_date)
    detail_query += " ORDER BY r.completion_date DESC, e.emp_no ASC"

    df_details = pd.read_sql_query(detail_query, conn, params=d_params)

    conn.close()

    output = io.BytesIO()
    with pd.ExcelWriter(output, engine='openpyxl') as writer:
        df_emp_summary.to_excel(writer, sheet_name='人員完訓時數總表', index=False)
        df_course_summary.to_excel(writer, sheet_name='課程開課統計表', index=False)
        df_details.to_excel(writer, sheet_name='完訓明細紀錄', index=False)

    output.seek(0)
    filename = f"護理之家_教育訓練簡易統計報表_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(output, download_name=filename, as_attachment=True)


if __name__ == '__main__':
    app.run(debug=True, port=5000)